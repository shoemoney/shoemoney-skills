"""Build a KDP-ready 6x9 DOCX: real styles, mirrored margins + gutter,
odd-page chapter starts, STYLEREF running heads, folios, auto TOC field."""
import json, os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from typo import smart, smart_runs

BOOK = json.load(open("build/book.json"))
M = BOOK["meta"]
BODY_FONT = "Georgia"          # present on every Word install, Mac and Windows
DISP_FONT = "Arial Narrow"     # ditto; the condensed display face

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e

def field(par, instr):
    r = par.add_run()
    r._r.append(el("w:fldChar", **{"w:fldCharType": "begin"}))
    t = el("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    r._r.append(t)
    r._r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    r._r.append(el("w:fldChar", **{"w:fldCharType": "end"}))
    return r

def setup_section(s, first=False):
    s.page_width, s.page_height = Inches(6), Inches(9)
    s.top_margin, s.bottom_margin = Inches(0.78), Inches(0.72)
    s.left_margin, s.right_margin = Inches(0.6875), Inches(0.6875)
    s.gutter = Inches(0.125)
    s.header_distance, s.footer_distance = Inches(0.42), Inches(0.4)
    s.different_first_page_header_footer = True
    return s

def style(doc, name, *, font=BODY_FONT, size=11.5, bold=False, italic=False,
          align=None, space_before=0, space_after=0, line=1.42, indent=0,
          caps=False, keep_with_next=False, page_break=False, base="Normal"):
    try:
        st = doc.styles[name]                   # builtin (Heading 1) -> modify in place
    except KeyError:
        st = doc.styles.add_style(name, 1)      # 1 = paragraph style
        if base != name:
            st.base_style = doc.styles[base]
    f = st.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = el("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), font)
    if caps:
        rpr.append(el("w:caps", **{"w:val": "1"}))
        rpr.append(el("w:spacing", **{"w:val": "30"}))
    p = st.paragraph_format
    if align is not None:
        p.alignment = align
    p.space_before, p.space_after = Pt(space_before), Pt(space_after)
    p.line_spacing = line
    p.first_line_indent = Inches(indent)
    p.keep_with_next = keep_with_next
    p.page_break_before = page_break
    p.widow_control = True
    return st

def build():
    doc = Document()
    n = doc.styles["Normal"]
    n.font.name, n.font.size = BODY_FONT, Pt(11.5)
    n.paragraph_format.line_spacing = 1.42
    n.paragraph_format.space_after = Pt(0)

    # mirrored margins live in settings.xml, not the section
    st = doc.settings.element
    st.append(el("w:mirrorMargins", **{"w:val": "1"}))
    st.append(el("w:evenAndOddHeaders", **{"w:val": "1"}))

    S = dict(
        BodyFirst=style(doc, "OCS Body First", align=WD_ALIGN_PARAGRAPH.JUSTIFY),
        Body=style(doc, "OCS Body", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.22),
        Lead=style(doc, "OCS Lead", bold=True, space_before=8, space_after=8),
        Sub=style(doc, "OCS Subhead", font=DISP_FONT, size=11, bold=True, caps=True,
                  space_before=16, space_after=5, keep_with_next=True),
        ChapNum=style(doc, "OCS Chapter Number", font=DISP_FONT, size=44, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER, space_before=54, space_after=2, line=1.0),
        ChapTitle=style(doc, "Heading 1", font=DISP_FONT, size=17, bold=True, caps=True,
                        align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=30,
                        line=1.25, keep_with_next=True, base="Heading 1"),
        FMHead=style(doc, "OCS Front Matter Head", font=DISP_FONT, size=15, bold=True, caps=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=44, space_after=24,
                     keep_with_next=True),
        Center=style(doc, "OCS Centered", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10),
        Fine=style(doc, "OCS Copyright", size=8.5, line=1.28, space_after=7),
        TitleBig=style(doc, "OCS Book Title", font=DISP_FONT, size=30, bold=True, caps=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER, line=1.05, space_after=0),
        TOCLine=style(doc, "OCS Contents", size=10.5, line=1.28, space_after=4),
        NotesHead=style(doc, "OCS Notes Head", font=DISP_FONT, size=14, bold=True, caps=True,
                        align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=6),
    )

    sec = setup_section(doc.sections[0], first=True)

    def para(style_name, text=None, runs=None, size=None):
        p = doc.add_paragraph(style=style_name)
        if runs is not None:
            for t, b, i in smart_runs(runs):
                for j, chunk in enumerate(t.split("\n")):
                    if j:
                        p.add_run().add_break()
                    r = p.add_run(chunk); r.bold = b; r.italic = i
        elif text is not None:
            for j, chunk in enumerate(smart(text).split("\n")):
                if j:
                    p.add_run().add_break()
                p.add_run(chunk)
        return p

    def pagebreak():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------- front matter ----------------
    para("OCS Book Title", M["title_line1"] + "\n" + M["title_line2"])
    pagebreak()
    para("OCS Book Title", M["title_line1"] + "\n" + M["title_line2"])
    p = para("OCS Centered", M["subtitle"]); p.runs[0].italic = True
    para("OCS Centered", "")
    para("OCS Centered", M["author"])
    pagebreak()
    for line in [l for l in BOOK["copyright"] if l.strip()]:
        para("OCS Copyright", line)
    pagebreak()
    para("OCS Front Matter Head", "Dedication")
    p = para("OCS Centered", BOOK["dedication"]); p.runs[0].italic = True
    pagebreak()

    para("OCS Front Matter Head", "Contents")
    tp = doc.add_paragraph(style="OCS Contents")
    field(tp, r' TOC \o "1-1" \h \z \u ')
    pagebreak()

    para("OCS Front Matter Head", "A Note of Gratitude")
    for i, b in enumerate(BOOK["gratitude"]):
        para("OCS Body First" if i == 0 else "OCS Body", runs=b["runs"])

    # ---------------- body: one odd-page section per chapter ----------------
    NOTE_PROMPTS = json.load(open("build/prompts.json"))

    def notes_page(label, prompt):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = doc.add_paragraph(style="OCS Notes Head"); p.add_run(smart(label))
        p.runs[0].font.size = Pt(8.5)
        p2 = doc.add_paragraph(style="OCS Notes Head"); p2.add_run("Your Notes")
        p3 = doc.add_paragraph(style="OCS Centered"); r = p3.add_run(smart(prompt))
        r.italic = True; r.font.size = Pt(9.5)
        for _ in range(18):
            lp = doc.add_paragraph()
            lp.paragraph_format.space_after = Pt(9)
            pb = lp._p.get_or_add_pPr()
            bd = el("w:pBdr"); bd.append(el("w:bottom", **{"w:val": "single", "w:sz": "6",
                                                          "w:space": "1", "w:color": "AAAAAA"}))
            pb.append(bd)

    def open_chapter():
        s = doc.add_section(WD_SECTION.ODD_PAGE)
        setup_section(s)
        s.header.is_linked_to_previous = True
        s.footer.is_linked_to_previous = True
        return s

    for c in BOOK["chapters"]:
        open_chapter()
        if c["num"]:
            para("OCS Chapter Number", c["num"])
        else:
            doc.add_paragraph(style="OCS Chapter Number")
        h = doc.add_paragraph(style="Heading 1"); h.add_run(smart(c["title"]))
        first = True
        for b in c["blocks"]:
            if b["kind"] == "subhead":
                doc.add_paragraph(style="OCS Subhead").add_run(
                    smart("".join(t for t, _, _ in b["runs"])))
                first = True
                continue
            sn = "OCS Lead" if b["kind"] == "lead" else \
                 ("OCS Body First" if first else "OCS Body")
            para(sn, runs=b["runs"])
            first = False
        idx = int(c["num"]) - 1 if c["num"] else None
        prompt = NOTE_PROMPTS[idx] if idx is not None and idx < len(NOTE_PROMPTS) else \
                 "What is one thing you will do differently tomorrow?"
        notes_page(f'Chapter {c["num"]} · {c["title"]}' if c["num"] else "Introduction", prompt)

    for key, title, blocks in (("onemore", "One More Thing", BOOK["one_more"]),
                               ("author", "About the Author", BOOK["author_bio"])):
        open_chapter()
        doc.add_paragraph(style="OCS Front Matter Head").add_run(smart(title))
        for i, b in enumerate(blocks):
            para("OCS Body First" if i == 0 else "OCS Body", runs=b["runs"])

    # ---------------- running heads + folios ----------------
    for i, s in enumerate(doc.sections):
        for hdr, is_even in ((s.header, False), (s.even_page_header, True)):
            hdr.is_linked_to_previous = False
            p = hdr.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.style = doc.styles["Header"]
            if i == 0:
                continue
            if is_even:
                r = p.add_run(M["title"])
            else:
                field(p, r' STYLEREF "Heading 1" \* MERGEFORMAT ')
                r = p.runs[-1]
            for run in p.runs:
                run.font.name, run.font.size, run.font.bold = DISP_FONT, Pt(8.5), True
                run.font.color.rgb = RGBColor(0, 0, 0)
                rpr = run._r.get_or_add_rPr()
                rpr.append(el("w:caps", **{"w:val": "1"}))
                rpr.append(el("w:spacing", **{"w:val": "26"}))
        for ftr in (s.footer, s.even_page_footer, s.first_page_footer):
            ftr.is_linked_to_previous = False
            p = ftr.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                continue
            field(p, " PAGE ")
            for run in p.runs:
                run.font.name, run.font.size = DISP_FONT, Pt(9)
                run.font.color.rgb = RGBColor(0, 0, 0)
        # chapter-opening pages: no running head
        s.first_page_header.is_linked_to_previous = False
        s.first_page_header.paragraphs[0].text = ""

    cp = doc.core_properties
    cp.title = "The One Call Salesperson: Stop Selling. Start Listening."
    cp.author = M["author"]
    cp.subject = M["subtitle"]
    cp.keywords = "new home sales, sales, homebuilding, listening, trust, sales training"
    cp.comments = ""
    out = "dist/OneCallSalesperson_Interior_6x9_KDP.docx"
    doc.save(out)
    print("wrote", out)

if __name__ == "__main__":
    build()
